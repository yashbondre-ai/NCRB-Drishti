from pathlib import Path
from docx import Document


def extract_docx(file_path: Path) -> str:
    document = Document(file_path)

    text = []

    # Paragraphs
    for para in document.paragraphs:
        if para.text.strip():
            text.append(para.text)

    # Tables
    for table in document.tables:
        for row in table.rows:
            row_data = []

            for cell in row.cells:
                if cell.text.strip():
                    row_data.append(cell.text.strip())

            if row_data:
                text.append(" | ".join(row_data))


    return "\n".join(text)